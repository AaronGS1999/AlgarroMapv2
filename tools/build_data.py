# -*- coding: utf-8 -*-
"""Curado y estandarizacion de coordenadas de las muestras de algarrobo.
Lee el Excel maestro y produce data/arboles.json + data/arboles.csv en WGS84 decimal.
Tambien enlaza cada arbol de Almeria con su ficha (carpeta Fichas)."""
import openpyxl, re, json, csv, unicodedata, os

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, "..", "Listado muestras ADN (final).xlsx")
OUT = os.path.join(HERE, "..", "data")
wb = openpyxl.load_workbook(XLSX, data_only=True)

corrections, warnings = [], []

def norm_code(v):
    if v is None: return None
    s = str(v).strip()
    if not s: return None
    if s.upper().startswith("CS"): s = s[2:].strip()
    try: return f"{int(float(s)):03d}"
    except Exception: return s

def strip_accents(s):
    if s is None: return ""
    return "".join(c for c in unicodedata.normalize("NFD", str(s)) if unicodedata.category(c) != "Mn")

def key(s):
    return re.sub(r"[^a-z0-9]", "", strip_accents(s).lower())

# ---------- parsers ----------
def to_num(t): return float(t.replace(",", "."))

def parse_decimal_pair(s):
    nums = re.findall(r"[-+]?\d+(?:[.,]\d+)?", s)
    if len(nums) < 2: return None
    return to_num(nums[0]), to_num(nums[1])

def parse_dms_token(tok, default_dir):
    if tok is None: return None
    t = (str(tok).replace("◦", "°").replace("º", "°")
         .replace("′", "'").replace("’", "'").replace("‘", "'")
         .replace("″", '"').replace("”", '"').replace("''", '"').replace("’’", '"'))
    d = default_dir
    m = re.search(r"[NSEWnsew]", t)
    if m: d = m.group(0).upper()
    t2 = re.sub(r"(?<=\d)\s+(?=\d)", "", t).replace(",", ".")
    parts = re.findall(r"\d+(?:\.\d+)?", t2)
    if not parts: return None
    val = float(parts[0]) + (float(parts[1])/60 if len(parts) > 1 else 0) + (float(parts[2])/3600 if len(parts) > 2 else 0)
    return -val if d in ("S", "W") else val

def looks_dms(s):
    return any(c in s for c in ["°", "◦", "º", "'", "′", "’"])

def split_dms_pair(s):
    m = re.search(r"[NSns]", s)
    if not m: return None
    lat = parse_dms_token(s[:m.end()], "N"); lon = parse_dms_token(s[m.end():], "W")
    return (lat, lon) if (lat is not None and lon is not None) else None

# ---------- Almeria ----------
alm = {}
for r in list(wb["Banco Almería"].iter_rows(values_only=True))[1:]:
    code = norm_code(r[12]); ubic = r[9]
    if not code or ubic is None: continue
    s = str(ubic).strip()
    if not s or strip_accents(s).lower() == "almeria": continue
    ll = split_dms_pair(s) if looks_dms(s) else None
    if ll is None: ll = parse_decimal_pair(s)
    if ll is None:
        warnings.append(f"Almeria {code}: no interpretable '{s}'"); continue
    lat, lon = ll
    if lon > 0:
        corrections.append(f"{code} ('{s}'): longitud positiva -> negativa"); lon = -lon
    if not (36.0 <= lat <= 38.0 and -3.6 <= lon <= -1.5):
        warnings.append(f"Almeria {code}: fuera de rango {lat:.5f},{lon:.5f}")
    alm[code] = (round(lat, 6), round(lon, 6))

# ---------- Marruecos ----------
hasna = {}
for r in list(wb["Datos Marruecos Hasna"].iter_rows(values_only=True))[1:]:
    if r[2] is None: continue
    try: sc = int(float(str(r[2]).strip()))
    except: continue
    hasna[sc] = dict(cod=norm_code(r[7]), lat=r[3], lon=r[4], site=r[1])

mor = {}
for r in list(wb["Marruecos_coord_correctas"].iter_rows(values_only=True))[1:]:
    if r[2] is None: continue
    try: sc = int(float(str(r[2]).strip()))
    except: continue
    lat = parse_dms_token(r[5], "N"); lon = parse_dms_token(r[6], "W")
    if lat is None or lon is None:
        warnings.append(f"Marruecos scode {sc}: DMS no interpretable"); continue
    mor[sc] = (round(lat, 6), round(lon, 6), "Marruecos_coord_correctas")

for sc, h in hasna.items():
    if sc in mor: continue
    lr, lonr = str(h["lat"]), str(h["lon"])
    if "X=" in lr or "Y=" in lr or "X=" in lonr or "Y=" in lonr:
        warnings.append(f"Marruecos scode {sc} ({h['site']}): solo UTM -> por region"); continue
    lat = parse_dms_token(lr, "N"); lon = parse_dms_token(lonr, "W")
    if lat is None or lon is None:
        warnings.append(f"Marruecos scode {sc}: respaldo no interpretable"); continue
    mor[sc] = (round(lat, 6), round(lon, 6), "Datos Marruecos Hasna (respaldo)")

mor_by_code = {}
for sc, (lat, lon, src) in mor.items():
    cod = hasna.get(sc, {}).get("cod") or f"{sc+515:03d}"
    if not (27.0 <= lat <= 37.0 and -13.5 <= lon <= -1.0):
        warnings.append(f"Marruecos {cod}: fuera de rango {lat:.4f},{lon:.4f}")
    mor_by_code[cod] = (lat, lon, src)

# ---------- gaceteer aproximado ----------
LOCALITIES = {
    "tavira (colecao antiga)": (37.127, -7.650), "tavira": (37.127, -7.650),
    "benafim": (37.242, -8.060), "paderne": (37.160, -8.204),
    "matos de cima, paderne": (37.170, -8.190), "querenca": (37.190, -7.995),
    "moncarapacho": (37.093, -7.783), "quelfes": (37.040, -7.833),
    "odeleite": (37.340, -7.462), "corte pequena, odeleite": (37.330, -7.470),
    "mesquita baixa, s b alportel": (37.222, -7.892),
    "varzea vinagre": (37.230, -8.040), "vila nova cacela": (37.163, -7.548),
    "selecao viveirista": (37.20, -8.05),
    "taroudant": (30.471, -8.877), "tarouant": (30.471, -8.877),
    "el ksiba": (32.573, -6.022), "essaouira": (31.513, -9.770),
    "moulay-idris el-methasiyne": (34.054, -5.522), "el menzel-beni asral": (33.836, -4.533),
    "talaint": (29.85, -9.30),
}
PROVINCES = {
    ("esp","tarragona"): (41.06, 1.10), ("esp","mallorca"): (39.62, 2.99),
    ("esp","malaga"): (36.72, -4.42), ("esp","alicante"): (38.35, -0.48),
    ("esp","castellon"): (40.10, -0.10), ("esp","valencia"): (39.47, -0.38),
    ("esp","ibiza"): (38.98, 1.43), ("esp","cadiz"): (36.53, -6.10),
    ("esp","barcelona"): (41.55, 1.95), ("esp","cordoba"): (37.89, -4.78),
    ("esp","sevilla"): (37.39, -5.99), ("esp","murcia"): (37.99, -1.30),
    ("esp","granada"): (37.18, -3.60), ("esp","almeria"): (37.10, -2.35),
    ("prt","algarve"): (37.20, -8.10),
    ("ita","sicilia"): (37.60, 14.02), ("ita","bari"): (41.12, 16.87),
    ("ita","cerdena"): (40.12, 9.01),
    ("aus","west"): (-31.95, 115.86), ("aus","geralpton"): (-28.774, 114.611),
    ("usa","california"): (36.78, -119.42),
    ("cyp","zona europea"): (35.00, 33.20),
    ("tur","sayini"): (39.00, 35.00), ("tur","de kaska"): (36.20, 29.64),
    ("mar","khenis-takot"): (33.30, -6.10), ("mar","essaouira"): (31.513, -9.770),
    ("mar","el ksiba"): (32.573, -6.022), ("mar","el menzel-beni asral"): (33.836, -4.533),
    ("mar","farsi"): (31.80, -7.10), ("mar","talaint"): (29.85, -9.30),
    ("mar","moulay-idris el-methasiyne"): (34.054, -5.522), ("mar","tarouant"): (30.471, -8.877),
    ("tun","ariana"): (36.86, 10.19), ("tun","bargou"): (36.09, 9.57),
    ("tun","sfax"): (34.74, 10.76), ("tun","jhadou"): (34.00, 9.00),
    ("tun","nomag"): (34.00, 9.00),
}
COUNTRIES = {
    "esp": (40.0, -3.7), "prt": (39.5, -8.0), "mar": (31.8, -7.09),
    "ita": (41.9, 12.5), "tun": (34.0, 9.5), "dza": (36.20, 3.60),
    "isr": (31.5, 34.9), "aus": (-31.95, 115.86), "usa": (37.5, -120.0),
    "hrv": (45.1, 15.2), "cyp": (35.0, 33.2), "tur": (39.0, 35.0),
}
CAJAMAR = (36.796, -2.720)

def resolve_region(pais, prov, origen, banco):
    pl = strip_accents(pais).lower().strip() if pais else ""
    pv = strip_accents(prov).lower().strip() if prov else ""
    og = strip_accents(origen).lower().strip() if origen else ""
    unk_pv = (pv in ("", "none")) or ("?" in pv)
    unk_og = (og in ("", "none")) or ("?" in og)
    if unk_pv and unk_og and banco == "CAJAMAR":
        return CAJAMAR, "Origen desconocido — conservado en CAJAMAR Las Palmerillas", "baja"
    if og in LOCALITIES:
        return LOCALITIES[og], f"Localidad: {origen}", "media"
    if (pl, pv) in PROVINCES:
        extra = f" ({origen})" if (og and og not in (pl, "none", "esp")) else ""
        return PROVINCES[(pl, pv)], f"Región: {prov}{extra}", "media"
    if pl in COUNTRIES:
        return COUNTRIES[pl], f"País: {pais}", "baja"
    return None, None, None

# ---------- fichas (accesion -> imagen, del proyecto v1) ----------
FICHAS_RAW = {
    "El Alquian": "Alquian.png", "Bedar 1": "Bedar1.png", "Bedar 2": "Bedar2.png",
    "Bedar 3": "Bedar3.png", "Belen 1": "Belen1.png", "Belen 2": "Belen2.png",
    "Botanico 1": "Botanico1.png", "Botanico 2": "Botanico2.png", "Botanico 3": "Botanico3.png",
    "Botanico 4": "Botanico4.png", "El algarrobico 1": "ElAlgarrobico1.png",
    "El algarrobico 2": "ElAlgarrobico2.png", "El algarrobico 3": "ElAlgarrobico3.png",
    "El de Cristobal": "ElDeCristobal.png", "El de Rafael": "ElDeRafael.png",
    "Enmedio 1": "Enmedio1.png", "Enmedio 2": "Enmedio2.png", "Enmedio 3": "Enmedio3.png",
    "Enmedio 4": "Enmedio4.png", "Enmedio 5": "Enmedio5.png", "Enmedio 6": "Enmedio6.png",
    "MacenasGolf": "GolfI.png", "Isleta 1": "Isleta1.png", "Isleta 2": "Isleta2.png",
    "La Serena 1": "LaSerena1.png", "La Serena 2": "LaSerena2.png",
    "Las Negras 1": "LasNegras1.png", "Las Negras 2": "LasNegras2.png", "Las Negras 3": "LasNegras3.png",
    "Las Niñas": "LasNinas.png", "Los Albacetes": "LosAlbacetes.png",
    "Los Mañas 1": "LosManas1.png", "Los Mañas 2": "LosManas2.png",
    "Lucainena 1": "Lucainena1.png", "Lucainena 2": "Lucainena2.png",
    "Lucainena 3": "Lucainena3.png", "Lucainena 4": "Lucainena4.png",
    "Marchalico 1": "Marchalico1.png", "Marchalico 2": "Marchalico2.png", "Marchalico 3": "Marchalico3.png",
    "Mojacar": "Mojacar.png", "MojacarAlto1": "MojacarAlto1.png", "MojacarAlto2": "MojacarAlto2.png",
    "MojacarAlto3": "MojacarAlto3.png", "MojacarAlto4": "MojacarAlto4.png",
    "MojacarAlto5": "MojacarAlto5.png", "MojacarAlto6": "MojacarAlto6.png",
    "Nijar": "Nijar.png", "Pileta 1": "Pileta1.png", "Pileta 2": "Pileta2.png",
    "Plomo 1": "Plomo1.png", "Plomo 2": "Plomo2.png", "Plomo 3": "Plomo3.png",
    "Rodalquilar": "Rodalquilar.png", "Rodena 1": "Rodena1.png", "Rodena 2": "Rodena2.png",
    "UAL1": "Ual1.png", "UAL2": "Ual2.png", "UAL3": "Ual3.png", "UAL4": "Ual4.png",
    "Vicar 1": "Vicar1.png", "Vicar 2": "Vicar2.png",
}
FICHAS = {key(k): v for k, v in FICHAS_RAW.items()}

# ---------- ensamblar ----------
def clean(v):
    if v is None: return None
    s = str(v).strip(); return s or None

def norm_wild(v):
    s = strip_accents(v).lower() if v else ""
    if not s: return "desconocido"
    if s.startswith("w") or "silvestre" in s: return "silvestre"
    if s.startswith("c") or "cultiv" in s: return "cultivado"
    if "actualmente no" in s: return "no cultivado actualmente"
    return "desconocido"

trees, seen, used_fichas = [], set(), set()
for r in list(wb["Listado común "].iter_rows(values_only=True))[1:]:
    code = norm_code(r[0])
    if not code or code in seen: continue
    seen.add(code)
    nombre = clean(r[2]); banco = clean(r[4]); origen = clean(r[5]); prov = clean(r[6]); pais = clean(r[8])
    lat = lon = precision = source = region = conf = None
    if code in alm:
        lat, lon = alm[code]; precision = "exacta"; source = "Banco Almería (GPS de campo)"
    elif code in mor_by_code:
        lat, lon, source = mor_by_code[code]; precision = "exacta"
    if lat is None:
        coords, region, conf = resolve_region(pais, prov, origen, banco)
        if coords:
            lat, lon = coords; precision = "aproximada"; source = "Ubicación representativa de la región de origen"
    ficha = FICHAS.get(key(nombre)) if nombre else None
    if ficha: used_fichas.add(ficha)
    trees.append(dict(id="Cs"+code, code=code, name=nombre, sex=clean(r[3]),
        bank=banco, origin=origen, province=prov, country=pais,
        wild=norm_wild(r[7]), wild_raw=clean(r[7]), lat=lat, lon=lon,
        precision=precision, coord_source=source, region_label=region,
        confidence=conf, ficha=ficha))

# ---------- salida ----------
import datetime
def keyf(t):
    try: return (0, int(t["code"]))
    except: return (1, t["code"])
trees.sort(key=keyf)
os.makedirs(OUT, exist_ok=True)
meta = dict(generated=datetime.date.today().isoformat(), source="Listado muestras ADN (final).xlsx",
            crs="WGS84 (EPSG:4326), grados decimales", count=len(trees),
            count_exact=sum(1 for t in trees if t["precision"] == "exacta"),
            count_approx=sum(1 for t in trees if t["precision"] == "aproximada"),
            count_fichas=sum(1 for t in trees if t["ficha"]))
with open(os.path.join(OUT, "arboles.json"), "w", encoding="utf-8") as f:
    json.dump(dict(meta=meta, trees=trees), f, ensure_ascii=False, indent=1)
cols = ["id","code","name","sex","bank","origin","province","country","wild","lat","lon","precision","coord_source","region_label","confidence","ficha"]
with open(os.path.join(OUT, "arboles.csv"), "w", encoding="utf-8", newline="") as f:
    w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore"); w.writeheader()
    for t in trees: w.writerow(t)

print(f"{len(trees)} árboles  ({meta['count_exact']} exactas, {meta['count_approx']} aproximadas)")
print(f"fichas enlazadas a {meta['count_fichas']} árboles; fichas usadas: {len(used_fichas)}/{len(FICHAS_RAW)}")
miss = [v for v in FICHAS_RAW.values() if v not in used_fichas]
if miss: print("fichas SIN emparejar:", miss)
print("correcciones:", len(corrections), "| avisos:", len(warnings))
